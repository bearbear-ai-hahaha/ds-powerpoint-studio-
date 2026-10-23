#!/usr/bin/env python3
"""Convert DS PowerPoint Studio .dspres JSON projects into real .pptx files."""
import argparse, json, sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

COLORS = {
    "navy": (10, 18, 45), "black": (0, 0, 0), "white": (255, 255, 255),
    "blue": (20, 70, 150), "green": (20, 105, 60), "red": (150, 35, 35),
    "purple": (95, 35, 130), "yellow": (245, 210, 40),
}

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input", help="input .dspres JSON file")
    ap.add_argument("output", help="output .pptx file")
    args = ap.parse_args()
    src = Path(args.input)
    data = json.loads(src.read_text(encoding="utf-8"))
    slides = data.get("slides", [])
    if not slides:
        raise SystemExit("Project has no slides.")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for item in slides:
        slide = prs.slides.add_slide(blank)
        color = COLORS.get(str(item.get("bg", "navy")).lower(), COLORS["navy"])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*color)
        title = str(item.get("title", "Untitled"))[:200]
        body = str(item.get("body", ""))[:4000]
        is_light = sum(color) > 400
        fg = RGBColor(20, 20, 25) if is_light else RGBColor(245, 246, 250)
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(11.9), Inches(1.35))
        tf = title_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Aptos Display"
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = fg
        body_box = slide.shapes.add_textbox(Inches(0.85), Inches(2.2), Inches(11.5), Inches(4.6))
        tf = body_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)
        tf.margin_top = Inches(0.1)
        tf.text = body
        for p in tf.paragraphs:
            p.font.name = "Aptos"
            p.font.size = Pt(22)
            p.font.color.rgb = fg
    out = Path(args.output)
    prs.save(out)
    print(f"Created {out} with {len(slides)} slides")

if __name__ == "__main__":
    main()
